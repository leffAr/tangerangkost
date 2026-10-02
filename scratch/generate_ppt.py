from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_ppt():
    prs = Presentation()

    # Define some slide layouts
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]
    title_only_layout = prs.slide_layouts[5]

    # --- Slide 1: Title ---
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "TangerangKost"
    subtitle.text = "Platform Direktori & Manajemen Kos Terintegrasi di Tangerang\n\nProyek Kelompok (6 Anggota)"

    # --- Slide 2: Latar Belakang ---
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "Latar Belakang & Masalah"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Pencari Kos: Kesulitan mencari kos valid dan sesuai budget di Tangerang tanpa harus survei langsung."
    p = tf.add_paragraph()
    p.text = "Pemilik Kos: Kesulitan memasarkan kos, mengelola kamar, dan berkomunikasi dengan penyewa."
    p = tf.add_paragraph()
    p.text = "Mobilitas: Mayoritas pencari kos mengakses info lewat HP (butuh aplikasi APK)."

    # --- Slide 3: Solusi ---
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "Solusi: TangerangKost"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Membangun platform berbasis Web & Aplikasi Android (APK) dengan sistem multi-peran (User, Owner, Admin)."
    p = tf.add_paragraph()
    p.text = "Keunggulan Utama:"
    p.level = 0
    p1 = tf.add_paragraph()
    p1.text = "Pencarian cerdas berbasis kecamatan se-Kabupaten Tangerang."
    p1.level = 1
    p2 = tf.add_paragraph()
    p2.text = "Sistem approval pemilik kos oleh Admin untuk menghindari penipuan."
    p2.level = 1
    p3 = tf.add_paragraph()
    p3.text = "Akses sangat cepat tanpa lag berkat optimasi Database Indexing."
    p3.level = 1

    # --- Slide 4: Jobdesc PM ---
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "1. Project Manager"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Apa yang Dikerjakan: Mengelola siklus proyek, membagi tugas, dan memastikan target mingguan tercapai."
    p = tf.add_paragraph()
    p.text = "Pengembangan Sistem: Menentukan prioritas fitur (Login -> Cari Kos -> Ulasan). Memastikan sinkronisasi Frontend & Backend."
    p = tf.add_paragraph()
    p.text = "Masalah & Solusi:"
    p1 = tf.add_paragraph()
    p1.text = "Masalah: Proses build APK Android memakan waktu dan error di komputer lokal."
    p1.level = 1
    p2 = tf.add_paragraph()
    p2.text = "Solusi: Menerapkan sistem GitHub Actions untuk otomatisasi build APK di Cloud."
    p2.level = 1

    # --- Slide 5: Jobdesc Analyst ---
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "2. Systems Analyst"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Apa yang Dikerjakan: Menganalisis kebutuhan pengguna, logika bisnis, dan mendesain skema Database (Relasi User, Kos, Kamar)."
    p = tf.add_paragraph()
    p.text = "Pengembangan Sistem: Merancang struktur approval Admin dan algoritma pencarian kecamatan."
    p = tf.add_paragraph()
    p.text = "Masalah & Solusi:"
    p1 = tf.add_paragraph()
    p1.text = "Masalah: Saat data bertambah, sistem pencarian kos melambat (lag) di server."
    p1.level = 1
    p2 = tf.add_paragraph()
    p2.text = "Solusi: Menganalisis query dan membuat Database Indexing (TiDB), pencarian turun ke milidetik."
    p2.level = 1

    # --- Slide 6: Jobdesc UI/UX ---
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "3. UI/UX Designer"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Apa yang Dikerjakan: Membuat desain UI/UX, menentukan palet warna (Biru), dan tata letak Dashboard."
    p = tf.add_paragraph()
    p.text = "Pengembangan Sistem: Menerapkan desain responsif untuk Laptop & HP. Merapikan desain font gradasi menjadi solid profesional."
    p = tf.add_paragraph()
    p.text = "Masalah & Solusi:"
    p1 = tf.add_paragraph()
    p1.text = "Masalah: Pengguna lupa kos mana yang sudah mereka lihat dan sukai."
    p1.level = 1
    p2 = tf.add_paragraph()
    p2.text = "Solusi: Merancang fitur 'Simpan Kos' (Ikon Hati) dan menu Sidebar 'Kos Tersimpan'."
    p2.level = 1

    # --- Slide 7: Jobdesc Frontend ---
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "4. Programmer (Frontend)"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Apa yang Dikerjakan: Koding Next.js (React) dan Tailwind CSS untuk antarmuka interaktif."
    p = tf.add_paragraph()
    p.text = "Pengembangan Sistem: Membangun pencarian real-time, rating dinamis, dan membungkus web menjadi APK dengan Capacitor."
    p = tf.add_paragraph()
    p.text = "Masalah & Solusi:"
    p1 = tf.add_paragraph()
    p1.text = "Masalah: Server GitHub gagal build APK (error path Node Modules & CRLF Windows)."
    p1.level = 1
    p2 = tf.add_paragraph()
    p2.text = "Solusi: Mengubah skrip build-apk.yml dengan pnpm exec dan inject skrip sed."
    p2.level = 1

    # --- Slide 8: Jobdesc Backend ---
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "5. Programmer (Backend)"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Apa yang Dikerjakan: Membangun logika otak aplikasi dengan NestJS dan TiDB (Prisma ORM)."
    p = tf.add_paragraph()
    p.text = "Pengembangan Sistem: Membangun REST API untuk Autentikasi (Bcrypt/JWT), upload gambar, dan CRUD data kos."
    p = tf.add_paragraph()
    p.text = "Masalah & Solusi:"
    p1 = tf.add_paragraph()
    p1.text = "Masalah: Gambar dari ImgBB diblokir oleh Internet Positif (Telkomsel/XL) di HP."
    p1.level = 1
    p2 = tf.add_paragraph()
    p2.text = "Solusi: Mengembangkan Image Proxy Edge Route untuk bypass sistem blokir secara rahasia via server Vercel."
    p2.level = 1

    # --- Slide 9: Jobdesc QA ---
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "6. Quality Assurance (Testing)"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Apa yang Dikerjakan: Menguji aplikasi sebagai Admin, Owner, dan User untuk mencari bug."
    p = tf.add_paragraph()
    p.text = "Pengembangan Sistem: Menguji instalasi file APK di HP Android asli dan melakukan stress-test loading halaman."
    p = tf.add_paragraph()
    p.text = "Masalah & Solusi:"
    p1 = tf.add_paragraph()
    p1.text = "Masalah: User tersesat di halaman profil karena tidak ada tombol kembali ke Dashboard, serta bug gambar error di paket data."
    p1.level = 1
    p2 = tf.add_paragraph()
    p2.text = "Solusi: Melaporkan temuan tersebut sehingga ditambahkan tombol 'Kembali ke Dashboard' dan memicu pembuatan Proxy."
    p2.level = 1

    # --- Slide 10: Penutup ---
    slide = prs.slides.add_slide(title_only_layout)
    title = slide.shapes.title
    title.text = "Kesimpulan & Selesai"
    
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = "TangerangKost berhasil dibangun berkat sinergi pembagian tugas 6 anggota yang terstruktur, menyelesaikan masalah dunia nyata secara efektif.\n\nTerima Kasih."
    p.alignment = PP_ALIGN.CENTER
    
    # Save presentation
    prs.save("TangerangKost_Presentasi.pptx")
    print("PPT created successfully!")

if __name__ == '__main__':
    create_ppt()
